import re
from collections import defaultdict
from datetime import datetime

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from ..kb import ROOT
from .content_quality import sanitize_generated_content
from .evidence_grading import canonical_url, classify as classify_evidence
from .evidence_review import usable as usable_evidence

EXPORT_DIR = ROOT / 'storage' / 'exports'
BLACK = RGBColor(0, 0, 0)
BLUE = RGBColor(22, 58, 87)
ACCENT = RGBColor(30, 104, 116)
MUTED = RGBColor(99, 108, 117)
LIGHT = 'EEF3F5'
BORDER = 'D6DEE2'


def _font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = 'Arial'
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn(f'w:{key}'), ('Microsoft YaHei' if bold else 'SimSun') if key == 'eastAsia' else 'Arial')
    if size is not None: run.font.size = Pt(size)
    if color is not None: run.font.color.rgb = color
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic


def _field(paragraph, instruction):
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    # Word fields belong inside runs, not directly inside w:p.
    for element in (begin, text, separate):
        paragraph.add_run()._r.append(element)
    paragraph.add_run('1')
    paragraph.add_run()._r.append(end)


def _shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd'); props.append(shading)
    shading.set(qn('w:fill'), fill)


def _cell_margins(cell, top=90, start=110, bottom=90, end=110):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar'); props.append(margins)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}'); margins.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')


def _table_borders(table, color=BORDER, size='4'):
    props = table._tbl.tblPr
    borders = props.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders'); props.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}'); borders.append(node)
        node.set(qn('w:val'), 'single'); node.set(qn('w:sz'), size); node.set(qn('w:color'), color)


def _mark_header(row):
    props = row._tr.get_or_add_trPr()
    header = props.find(qn('w:tblHeader'))
    if header is None:
        header = OxmlElement('w:tblHeader'); props.append(header)
    header.set(qn('w:val'), 'true')


def _styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    normal.font.size = Pt(11); normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.widow_control = True
    specs = [('Title', 29, 0, 14), ('Subtitle', 12, 0, 24),
             ('Heading 1', 18, 22, 12), ('Heading 2', 13, 14, 6), ('Heading 3', 11, 10, 4)]
    specs.extend((f'Heading {level}', 11, 8, 4) for level in range(4, 10))
    for name, size, before, after in specs:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'; style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(size); style.font.bold = name != 'Subtitle'; style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ('List Bullet', 'List Number'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'; style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(10.2); style.font.color.rgb = BLACK
        style.paragraph_format.left_indent = Cm(.65); style.paragraph_format.first_line_indent = Cm(-.3)
        style.paragraph_format.space_after = Pt(4); style.paragraph_format.line_spacing = 1.35


def _header_footer(section, title):
    section.header_distance = Mm(10); section.footer_distance = Mm(10)
    hp = section.header.paragraphs[0]
    _font(hp.add_run('CKOS RESEARCH'), 8, MUTED, True); hp.add_run('    '); _font(hp.add_run(title), 8, MUTED)
    fp = section.footer.paragraphs[0]
    _font(fp.add_run(datetime.now().strftime('%Y.%m.%d')), 8, MUTED); fp.add_run('\t')
    _font(fp.add_run('行业研究报告    '), 8, MUTED); _field(fp, 'PAGE')
    fp.paragraph_format.tab_stops.add_tab_stop(Mm(165))


def _cover(doc, project):
    brief = project['brief']
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(72); p.paragraph_format.space_after = Pt(36)
    _font(p.add_run('CKOS  /  INDUSTRY RESEARCH'), 9, ACCENT, True)
    title = doc.add_paragraph(style='Title'); _font(title.add_run(project['title']), 29, BLACK, True)
    subtitle = doc.add_paragraph(style='Subtitle'); _font(subtitle.add_run('行业全景、竞争格局与未来机会'), 12, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(72)
    rows = [('研究地区', brief.get('geography', '')),
            ('研究期间', f"{brief.get('history_start')}—{brief.get('history_end')}  |  预测至 {brief.get('forecast_end')}"),
            ('研究目的', brief.get('purpose', '')), ('报告日期', datetime.now().strftime('%Y年%m月%d日'))]
    table = doc.add_table(rows=len(rows), cols=2); table.alignment = WD_TABLE_ALIGNMENT.LEFT; table.autofit = False
    _mark_header(table.rows[0])
    table.columns[0].width = Mm(27); table.columns[1].width = Mm(118); _table_borders(table, 'FFFFFF', '0')
    for row, (label, value) in zip(table.rows, rows):
        for cell in row.cells: _cell_margins(cell, 70, 0, 70, 0)
        _font(row.cells[0].paragraphs[0].add_run(label), 9, MUTED, True)
        _font(row.cells[1].paragraphs[0].add_run(str(value)), 10, BLACK)
    brand = doc.add_paragraph(); brand.paragraph_format.space_before = Pt(76)
    _font(brand.add_run('企业研究工作台'), 9, ACCENT, True)
    doc.add_page_break()


def _toc(doc, chapters):
    doc.add_heading('目录', level=1)
    for chapter in chapters:
        label = '摘要' if chapter['chapter_no'] == 0 else f"{chapter['chapter_no']:02d}"
        table = doc.add_table(rows=1, cols=2); table.autofit = False
        _mark_header(table.rows[0])
        table.columns[0].width = Mm(18); table.columns[1].width = Mm(132); _table_borders(table, 'FFFFFF', '0')
        for cell in table.rows[0].cells: _cell_margins(cell, 70, 0, 70, 0)
        _font(table.cell(0, 0).paragraphs[0].add_run(label), 9, ACCENT, True)
        _font(table.cell(0, 1).paragraphs[0].add_run(chapter['title']), 10.5, BLACK, True)
    doc.add_page_break()


def _inline(paragraph, text, size=11, color=BLACK):
    for part in re.split(r'(\*\*.*?\*\*)', text):
        if not part: continue
        bold = part.startswith('**') and part.endswith('**')
        for token in re.split(r'(\[\d+(?:-\d+)?\])', part[2:-2] if bold else part):
            run = paragraph.add_run(token)
            citation = bool(re.fullmatch(r'\[\d+(?:-\d+)?\]', token))
            _font(run, 8 if citation else size, color, bold)
            if citation: run.font.superscript = True


def _numbered(doc, text, start, num_id=None):
    root = doc.part.numbering_part.element
    if num_id is None:
        abstract_id = max([int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))] + [-1]) + 1
        abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(abstract_id))
        level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
        for name, value in [('start', str(start)), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('suff', 'tab')]:
            node = OxmlElement('w:' + name); node.set(qn('w:val'), value); level.append(node)
        abstract.append(level); root.append(abstract)
        num_id = max([int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))] + [0]) + 1
        num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
        ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(abstract_id)); num.append(ref); root.append(num)
    p = doc.add_paragraph(style='List Number')
    props = p._p.get_or_add_pPr(); numbering = OxmlElement('w:numPr')
    for name, value in [('ilvl', 0), ('numId', num_id)]:
        node = OxmlElement('w:' + name); node.set(qn('w:val'), str(value)); numbering.append(node)
    props.append(numbering); p.paragraph_format.tab_stops.add_tab_stop(Cm(.65))
    _inline(p, text)
    return num_id


def _normalized_heading(text):
    return re.sub(r'[\s·:：一二三四五六七八九十第章节、，。_-]+', '', text).lower()


def _clean_markdown(content, outer_title):
    lines = content.splitlines(); outer = _normalized_heading(outer_title)
    for index, raw in enumerate(lines):
        line = raw.strip()
        if not line: continue
        match = re.match(r'^#{1,6}\s*(.+?)\s*#*$', line)
        if match:
            inner = _normalized_heading(match.group(1))
            if inner and (inner == outer or inner in outer or outer in inner or re.match(r'^第?\d+章', match.group(1))):
                lines.pop(index)
        break
    return '\n'.join(lines)


def _markdown_table(doc, rows):
    parsed = [[part.strip() for part in row.strip().strip('|').split('|')] for row in rows]
    if len(parsed) > 1 and all(re.fullmatch(r':?-{3,}:?', value.replace(' ', '')) for value in parsed[1]):
        parsed.pop(1)
    width = max(len(row) for row in parsed); parsed = [row + [''] * (width - len(row)) for row in parsed]
    table = doc.add_table(rows=len(parsed), cols=width); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _mark_header(table.rows[0])
    table.autofit = True; _table_borders(table)
    for r_idx, values in enumerate(parsed):
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_margins(cell)
            if r_idx == 0: _shade(cell, '163A57')
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
            _inline(p, value, 8.5, RGBColor(255, 255, 255) if r_idx == 0 else BLACK)
            if r_idx == 0:
                for run in p.runs: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _heading_label(text):
    # Only heading numbering; preserve titles such as 2026年展望 and 3D技术.
    return re.sub(r'^(?:\d{1,2}(?:\.\d{1,2})+(?:[.、)]?\s*|(?=[\u4e00-\u9fff]))|\d{1,2}[.、)]\s+|[一二三四五六七八九十]+[、．]\s*)', '', text).strip()


def _markdown(doc, content, outer_title='', chapter_no=None):
    content = sanitize_generated_content(content)
    if chapter_no and chapter_no > 0:
        content = re.sub(r'\[(\d+)\]', lambda match: f'[{chapter_no}-{match.group(1)}]', content)
    content = _clean_markdown(content, outer_title) if outer_title else content
    lines = content.splitlines(); index = 0; num_id = None
    heading_stack = []
    heading_counts = []
    generic = {'结论', '核心观点', '事实与原因', '影响', '限制说明', '证据不足'}
    while index < len(lines):
        line = lines[index].strip()
        if not line: index += 1; continue
        ordered = re.match(r'^(\d+)[.)、]\s*(.+)', line)
        if not ordered: num_id = None
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip()); index += 1
            if len(rows) >= 2: _markdown_table(doc, rows)
            else:
                p = doc.add_paragraph(); _inline(p, rows[0].strip('|').replace('|', '  /  '))
            continue
        heading = re.match(r'^#{1,6}\s*(.+?)\s*#*$', line)
        if heading:
            text = re.sub(r'\*\*', '', heading.group(1)).strip()
            if chapter_no is not None and chapter_no > 0:
                depth = len(line) - len(line.lstrip('#'))
                # Relative hierarchy: ### can be a main section after the outer
                # chapter heading is removed; #### remains its child.
                while heading_stack and heading_stack[-1] > depth:
                    heading_stack.pop()
                if not heading_stack or heading_stack[-1] < depth:
                    heading_stack.append(depth)
                level = len(heading_stack)
                heading_counts = heading_counts[:level]
                if len(heading_counts) < level:
                    heading_counts.append(0)
                heading_counts[-1] += 1
                label = '.'.join(map(str, [chapter_no] + heading_counts))
                doc.add_heading(f'{label} {_heading_label(text)}', level=min(9, level + 1))
            elif text in generic:
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
                _font(p.add_run(text), 11, BLACK, True)
            else: doc.add_heading(text, level=min(3, max(2, len(line) - len(line.lstrip('#')))))
        elif re.match(r'^[-*+]\s+', line):
            p = doc.add_paragraph(style='List Bullet'); _inline(p, re.sub(r'^[-*+]\s+', '', line))
        elif ordered:
            number = int(ordered.group(1))
            if number == 1: num_id = None
            num_id = _numbered(doc, ordered.group(2), number, num_id)
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(.74); _inline(p, line)
        index += 1


def _research_scope(doc, brief):
    doc.add_heading('研究说明', level=1)
    rows = [('研究范围', brief.get('included_segments') or '以研究任务书所界定行业为准'),
            ('排除范围', brief.get('excluded_segments') or '未特别指定'),
            ('重点关注', brief.get('focus') or '未特别指定'),
            ('重点企业', brief.get('key_companies') or '未特别指定')]
    table = doc.add_table(rows=len(rows), cols=2); table.autofit = False
    _mark_header(table.rows[0])
    table.columns[0].width = Mm(28); table.columns[1].width = Mm(122); _table_borders(table)
    for idx, (label, value) in enumerate(rows):
        left, right = table.rows[idx].cells; _shade(left, LIGHT)
        for cell in (left, right): _cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _font(left.paragraphs[0].add_run(label), 9, BLUE, True); _font(right.paragraphs[0].add_run(str(value)), 9.5, BLACK)
    doc.add_page_break()


def _reference_material(chapters, evidence):
    """Move explicit bibliography entries, not analytical source/limitations prose."""
    cleaned, entries = [], []
    for chapter in chapters:
        lines = (chapter.get('content') or '').splitlines(); kept = []; in_refs = False
        for line in lines:
            heading = re.match(r'^\s*#{1,6}\s*(.*?)\s*$', line)
            if heading:
                name = heading.group(1).replace('**', '').strip(' ：:')
                in_refs = name in {'参考文献', '引用文献', '本章证据', '本章参考文献', '参考资料'}
                if in_refs: continue
            match = re.match(r'^\s*(?:[-*]\s*)?\[(\d+)\]\s*(.+)', line)
            if in_refs and match:
                entries.append(dict(title=match.group(2), chapter_no=chapter['chapter_no'], local_id=match.group(1), url=''))
            elif line.strip():
                kept.append(line)
        cleaned.append(dict(chapter, content='\n'.join(kept)))
    # Match explicit author-supplied lists by title, not a guessed renumbering.
    for entry in entries:
        title = entry['title'].split('｜')[0].strip()
        matched = next((e for e in evidence if e.get('chapter_no') == entry['chapter_no'] and (e.get('title') or '').strip() == title), None)
        if matched:
            entry.update(title=matched['title'], url=matched.get('url', ''), publisher=matched.get('publisher', ''), published_at=matched.get('published_at', ''))
    return cleaned, entries


def _cited_ids_by_chapter(chapters):
    """Collect local evidence numbers that are actually cited in chapter prose.

    A compact ``[2-4]`` citation is treated as the inclusive local range.  In
    the executive summary, ``[4-2]`` is interpreted as chapter 4, source 2;
    bare summary references cannot be mapped safely and are therefore ignored.
    """
    cited = defaultdict(set)
    for chapter in chapters:
        chapter_no = int(chapter.get('chapter_no') or 0)
        for match in re.finditer(r'\[(\d+)(?:-(\d+))?\]', chapter.get('content') or ''):
            first, second = int(match.group(1)), match.group(2)
            if chapter_no == 0:
                if second is not None and 1 <= first <= 10:
                    cited[first].add(int(second))
                continue
            if second is None:
                cited[chapter_no].add(first)
                continue
            last = int(second)
            # Citation ranges are intentionally bounded; a bracketed year span
            # such as [2021-2025] must never create thousands of references.
            if first <= last and last - first <= 20:
                cited[chapter_no].update(range(first, last + 1))
    return {chapter: sorted(ids) for chapter, ids in cited.items()}


def _sources(doc, evidence, explicit=None, cited_by_chapter=None):
    """Render cited sources with the exact chapter-local numbers used in prose.

    ``cited_by_chapter=None`` keeps the small internal helper backward
    compatible.  Normal report export always supplies the citation map and
    therefore omits collected-but-unused material.
    """
    unique = {}

    def add(item, chapter, local_id):
        key = canonical_url(item.get('url') or '') or (item.get('title') or '').strip()
        if not key:
            return
        entry = unique.setdefault(key, dict(item, chapters=set(), refs=set()))
        entry['chapters'].add(chapter)
        if local_id is not None:
            entry['refs'].add(f"{chapter}-{local_id}")

    evidence = list(evidence)
    explicit = list(explicit or [])
    if cited_by_chapter is None:
        counts = defaultdict(int)
        for item in evidence:
            chapter = item.get('chapter_no')
            counts[chapter] += 1
            add(item, chapter, counts[chapter])
        for item in explicit:
            chapter = item.get('chapter_no')
            add(item, chapter, item.get('local_id'))
    else:
        grouped = defaultdict(list)
        for item in evidence:
            grouped[item.get('chapter_no')].append(item)
        explicit_by_ref = {
            (item.get('chapter_no'), int(item['local_id'])): item
            for item in explicit if str(item.get('local_id') or '').isdigit()
        }
        for chapter, local_ids in sorted(cited_by_chapter.items()):
            rows = grouped.get(chapter, [])
            for local_id in local_ids:
                item = rows[local_id - 1] if 0 < local_id <= len(rows) else explicit_by_ref.get((chapter, local_id))
                if item:
                    add(item, chapter, local_id)
    if not unique: return
    doc.add_page_break(); doc.add_heading('参考资料', level=1)
    p = doc.add_paragraph('相同原始来源合并列示；章内参考编号以“章号-编号”保留。A—D为来源类型辅助分级，不代替指标口径核验。')
    _font(p.runs[0], 9, MUTED)
    for entry in unique.values():
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_together = True
        refs = '、'.join(sorted(entry['refs']))
        prefix = f'[{refs}] ' if refs else '〔第'+ '、'.join(str(x) for x in sorted(entry['chapters'], key=lambda x:x or 0))+'章资料〕 '
        grade = classify_evidence(entry)['grade']
        _font(p.add_run(prefix + f'[{grade}] ' + entry.get('title', '')), 9, BLACK)
        meta = ' · '.join(str(x) for x in [entry.get('publisher'), (entry.get('published_at') or '')[:10]] if x)
        if meta: _font(p.add_run('  '+meta), 8.5, MUTED)
        url = entry.get('url') or ''
        if url.startswith(('https://', 'http://')):
            link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = p.add_run('  原文'); _font(run, 9, ACCENT); run.font.underline = True
            link.append(run._r); p._p.append(link)


def _chapter_dashboard(doc, chart_specs, start=1):
    if not chart_specs: return
    for index, spec in enumerate(chart_specs, start):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        _font(p.add_run(f"图 {spec['chapter_no']}-{index}  {spec['title']}"), 10.5, BLACK, True)
        picture = doc.add_paragraph(); picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.keep_with_next = True
        shape = picture.add_run().add_picture(str(spec['image_path']), width=Cm(14.8))
        doc_props = shape._inline.docPr
        doc_props.set('title', spec['title'])
        doc_props.set('descr', f"{spec['title']}，单位{spec['unit']}，数据来自第{spec['chapter_no']}章")
        source_text = (f"单位：{spec['unit']}  |  数据来自结构化决策记录及程序计算"
                       if str(spec.get('kind') or '').startswith('decision_')
                       else f"单位：{spec['unit']}  |  数据提取自第{spec['chapter_no']}章正文")
        source = doc.add_paragraph(source_text)
        source.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _font(source.runs[0], 7.5, MUTED)
        if spec.get('note'):
            note = doc.add_paragraph(spec['note']); _font(note.runs[0], 8, MUTED)


def _anchor_text(text):
    # Markdown list markers are not present in Word paragraph text.
    # Strip only leading markers, preserving negative values and numeric data.
    text = re.sub(r'^\s*(?:[-*+]\s+|\d+[.)、]\s*)', '', text)
    text = re.sub(r'\[\d+(?:-\d+)?\]', '', text)
    return re.sub(r'[\s#*_`|]+', '', text)


def _place_chapter_charts(doc, blocks, chart_specs):
    """Exact source-excerpt anchoring only, scoped to rendered chapter blocks.

    A series follows its last source block. Missing or ambiguous source excerpts
    fall back to the chapter end instead of guessing from a shared number.
    """
    candidates = []
    for index, block in enumerate(blocks):
        style = block.find('./' + qn('w:pPr') + '/' + qn('w:pStyle'))
        if style is not None and (style.get(qn('w:val')) or '').startswith('Heading'):
            continue
        value = ''.join(block.xpath('.//w:t/text()'))
        candidates.append((index, _anchor_text(value)))
    placements = []
    for spec in chart_specs:
        excerpts = list(dict.fromkeys(_anchor_text(x) for x in spec.get('source_excerpts', []) if _anchor_text(x)))
        positions = []
        for excerpt in excerpts:
            matches = [i for i, text in candidates if len(excerpt) >= 6 and excerpt in text]
            if len(matches) != 1:
                positions = []
                break
            positions.append(matches[0])
        placements.append((max(positions) if positions else len(blocks), spec))
    ordered = sorted(placements, key=lambda item: item[0])
    references = defaultdict(list)
    for number, (position, spec) in enumerate(ordered, 1):
        if position < len(blocks) and spec.get('chapter_no') is not None:
            references[position].append(f"图 {spec['chapter_no']}-{number}")
    tails = {}
    for position, labels in references.items():
        block = blocks[position]
        existing_text = ''.join(block.xpath('.//w:t/text()'))
        missing = [label for label in labels if not re.search(re.escape(label).replace(r'\ ', r'\s*') + r'(?!\d)', existing_text)]
        if not missing:
            continue
        if block.tag == qn('w:p'):
            paragraph = Paragraph(block, doc._body)
            prefix = '' if not paragraph.text or paragraph.text.rstrip().endswith(('。','！','？','；')) else '。'
        else:
            # A table anchor gets its own lead-in, not text stuffed into a cell.
            paragraph = doc.add_paragraph()
            block.addnext(paragraph._p)
            tails[position] = paragraph._p
            prefix = ''
        _inline(paragraph, prefix + '相关数据如' + '、'.join(missing) + '所示。')
    for number, (position, spec) in enumerate(ordered, 1):
        if position >= len(blocks):
            lead = doc.add_paragraph()
            lead.paragraph_format.space_before = Pt(10)
            lead.paragraph_format.space_after = Pt(5)
            figure_label = (f"图 {spec['chapter_no']}-{number}"
                            if spec.get('chapter_no') is not None else f"图 {number}")
            _inline(lead, f"本章相关测算或程序测算结果如{figure_label}所示。")
        existing = set(doc._element.body)
        _chapter_dashboard(doc, [spec], start=number)
        added = [node for node in doc._element.body if node not in existing]
        if position < len(blocks):
            anchor = tails.get(position, blocks[position])
            for node in added:
                anchor.addnext(node)
                anchor = node
            tails[position] = anchor


def build_docx(project, evidence, chart_specs=None, *, export_id=None):
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    if export_id is not None and not re.fullmatch(r'[a-f0-9]{32}', export_id):
        raise ValueError('Invalid export id')
    path = EXPORT_DIR / f'{export_id or project["id"]}.docx'
    # Use the same evidence eligibility rule as generation and project audits.
    # This defensive filter also protects historical exports and direct callers.
    evidence = usable_evidence(evidence)
    doc = Document(); section = doc.sections[0]
    section.page_width = Mm(210); section.page_height = Mm(297)
    section.top_margin = Mm(22); section.bottom_margin = Mm(20); section.left_margin = Mm(27); section.right_margin = Mm(23)
    _styles(doc); _header_footer(section, project['title']); _cover(doc, project)
    chapters = sorted(project['chapters'], key=lambda x: x['chapter_no'])
    chapters, explicit_refs = _reference_material(chapters, evidence)
    _toc(doc, chapters); _research_scope(doc, project['brief'])
    charts_by_chapter = defaultdict(list)
    for spec in chart_specs or []: charts_by_chapter[spec['chapter_no']].append(spec)
    for pos, chapter in enumerate(chapters):
        number = 'EXECUTIVE SUMMARY' if chapter['chapter_no'] == 0 else f"CHAPTER {chapter['chapter_no']:02d}"
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4); _font(p.add_run(number), 8.5, ACCENT, True)
        heading = '摘要：核心结论' if chapter['chapter_no'] == 0 else f"第{chapter['chapter_no']}章  {chapter['title']}"
        doc.add_heading(heading, level=1)
        existing = set(doc._element.body)
        _markdown(doc, chapter.get('content') or '本章尚未生成。', chapter['title'], chapter['chapter_no'])
        blocks = [node for node in doc._element.body if node not in existing]
        _place_chapter_charts(doc, blocks, charts_by_chapter.get(chapter['chapter_no'], []))
        if pos < len(chapters) - 1: doc.add_page_break()
    cited = _cited_ids_by_chapter(chapters)
    if cited: _sources(doc, evidence, explicit_refs, cited)
    props = doc.core_properties; props.title = project['title']; props.author = 'CKOS 企业研究工作台'; props.subject = '行业研究报告'
    props.comments = '由 CKOS 行业研究系统生成'; doc.save(path)
    return path
