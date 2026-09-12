"""Read-only inspection of report paragraphs, captions and repeated sentences."""
import sys,re,json
from collections import defaultdict
from docx import Document
doc=Document(sys.argv[1])
occurrences=defaultdict(list)
chapter='前置内容'
captions=[];references=[];samples=[]
for index,p in enumerate(doc.paragraphs):
    text=p.text.strip()
    if p.style.name=='Heading 1':chapter=text
    if re.match(r'^图\s*\d',text):captions.append((index,chapter,text))
    if re.search(r'如图|如\s*图|见图',text):references.append((index,text))
    if text and len(samples)<20 and chapter.startswith('第'):samples.append((index,text))
    if chapter.startswith('参考'):continue
    for sentence in re.split(r'(?<=[。！？])',text):
        normalized=re.sub(r'\[\d+(?:-\d+)?\]|\s','',sentence)
        if len(normalized)>25:occurrences[normalized].append((index,chapter))
duplicates=[dict(text=k,locations=v) for k,v in occurrences.items() if len(v)>1]
print(json.dumps(dict(paragraphs=len(doc.paragraphs),images=len(doc.inline_shapes),captions=captions,references=references,duplicates=duplicates[:30],samples=samples),ensure_ascii=False,indent=2))
