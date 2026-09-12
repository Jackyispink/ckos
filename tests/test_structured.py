from docx import Document

from app import database, kb, structured


def test_scoped_diagnostics_and_assets(tmp_path, monkeypatch):
    monkeypatch.setattr(database, 'ROOT', tmp_path)
    monkeypatch.setattr(kb, 'ROOT', tmp_path)
    (tmp_path / 'data').mkdir()
    for name in ('甲甲有限公司', '乙乙有限公司'):
        d = Document()
        d.add_paragraph(name)
        d.add_heading('人力资源管理诊断与建议', level=2)
        d.add_paragraph('2. 存在现状')
        d.add_paragraph(name + '缺少人才梯队。')
        d.add_paragraph('3. 改善建议')
        d.add_paragraph(name + '应制定后备人才培养计划。')
        d.save(tmp_path / 'data' / (name + '管理诊断报告.docx'))
    kb.ingest()
    assert structured.ingest(tmp_path)['processed'] == 2
    assert structured.ingest(tmp_path)['unchanged'] == 2
    sources, _ = structured.resolve('甲甲有限公司诊断建议是什么')
    assert len(sources) == 1
    assert '乙乙' not in sources[0]['text']
    assert sources[0]['evidence']
    followup, _ = structured.resolve('这家企业的人才怎么改', previous_sources=sources)
    assert followup[0]['company_id'] == sources[0]['company_id']
    switched, _ = structured.resolve('乙乙有限公司诊断建议', previous_sources=sources)
    assert switched[0]['company_id'] != sources[0]['company_id']
    assert structured.resolve('未知企业的诊断建议')[0] == []
    assert structured.resolve('乙乙有限公司诊断建议', document_id=sources[0]['document_id'])[0] == []
    with database.connect() as db:
        rel = db.execute('SELECT * FROM kb_relations LIMIT 1').fetchone()
        assert rel['review_status'] == 'pending'
        assert rel['basis'] == 'same_module_candidate'


def test_table_cells_not_deduplicated(tmp_path):
    d = Document()
    d.add_paragraph('测试有限公司')
    table = d.add_table(rows=2, cols=3)
    for c, text in zip(table.rows[0].cells, ['指标', '满分', '得分']):
        c.text = text
    for c, text in zip(table.rows[1].cells, ['指标A', '5', '5']):
        c.text = text
    path = tmp_path / '测试有限公司.docx'
    d.save(path)
    _, _, blocks, _ = structured.parse(path)
    row = [b for b in blocks if b['kind'] == 'table_row'][-1]
    assert [c['text'] for c in row['metadata']['cells']] == ['指标A', '5', '5']


def test_conflicting_scores_preserved(tmp_path):
    d = Document()
    d.add_paragraph('测试有限公司')
    table = d.add_table(rows=2, cols=3)
    for c, text in zip(table.rows[0].cells, ['维度', '分值', '得分']):
        c.text = text
    for c, text in zip(table.rows[1].cells, ['营销管理', '130', '86.5']):
        c.text = text
    d.add_heading('营销管理诊断与建议', level=2)
    d.add_paragraph('营销管理模块满分为100分，得分86.5分。')
    path = tmp_path / '测试有限公司.docx'
    d.save(path)
    _, _, _, quality = structured.parse(path)
    assert quality['score_conflicts'][0]['table_value'] == 130
    assert quality['score_conflicts'][0]['text_value'] == 100
