from io import BytesIO
import json
import time
from threading import Event

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import database,kb,upload_preview,duplicates,duplicate_index,upload_jobs,uploads
from app.main import app


def word(company='测试科技有限公司'):
    d=Document();d.add_paragraph(company);d.add_paragraph('管理诊断报告');d.add_paragraph('改善设备管理流程与人才培养。')
    out=BytesIO();d.save(out);return out.getvalue()


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(kb,'ROOT',tmp_path);monkeypatch.setattr(database,'ROOT',tmp_path)
    with TestClient(app) as client:yield client


def preview(client,payload=None,name='诊断报告.docx'):
    return client.post('/api/uploads/preview',files={'file':(name,payload or word())})


def test_duplicate_previews_preserved_without_database_checks(client,monkeypatch):
    monkeypatch.setattr(duplicates,'check_names',lambda *a,**k:pytest.fail('no name rejection before review'))
    monkeypatch.setattr(duplicate_index,'check',lambda *a,**k:pytest.fail('no similarity computation before review'))
    monkeypatch.setattr(kb,'ingest',lambda *a,**k:pytest.fail('no chunking before review'))
    first=preview(client).json();second=preview(client).json()
    assert first['token']!=second['token']
    assert first['company_name']==second['company_name']=='测试科技有限公司'
    assert first['report_title']==second['report_title']=='诊断报告'
    assert client.get('/api/status').json()['documents']==[]
    with database.connect() as db:
        for table in ['documents','kb_reports','chunks','dedup_documents','dedup_models']:
            assert db.execute(f'SELECT count(*) n FROM {table}').fetchone()['n']==0


def test_manual_names_override_and_commit_only_after_review(client):
    info=preview(client).json()
    body={'token':info['token'],'company_name':'人工确认有限公司','report_title':'2026年管理报告'}
    response=client.post('/api/uploads/commit',json=body)
    assert response.status_code==201,response.text
    assert client.get('/api/manage/reports').json()[0]['company']=='人工确认有限公司'
    assert not (upload_preview.folder()/(info['token']+'.docx')).exists()
    assert client.post('/api/uploads/commit',json=body).status_code==404


def test_rejected_commit_keeps_staging_for_correction(client):
    original=preview(client).json()
    body={'token':original['token'],'company_name':'测试科技有限公司','report_title':'年度报告'}
    assert client.post('/api/uploads/commit',json=body).status_code==201
    another=preview(client).json();body['token']=another['token']
    result=client.post('/api/uploads/commit',json=body)
    assert result.status_code==409 and result.json()['detail']['code']=='name_conflict'
    assert upload_preview.resolve(another['token'])[0].exists()


def test_unknown_and_conflicting_company_names():
    assert upload_preview.suggest('扫描件1.docx',['没有企业名称的报告'])['needs_review']
    r=upload_preview.suggest('01甲甲有限公司管理诊断报告.docx',['乙乙有限公司','管理诊断报告'])
    assert r['needs_review'] and r['candidates']==['乙乙有限公司','甲甲有限公司']
    r=upload_preview.suggest('扫描件1.docx',['企业名称：中稀（常熟）稀土新材料有限公司','管理诊断报告'])
    assert r['company_name']=='中稀(常熟)稀土新材料有限公司'
    assert '管理诊断报告' in r['report_title']


def test_invalid_expired_staging(client):
    assert preview(client,b'not a docx').status_code==422
    assert not list(upload_preview.folder().glob('*'))
    assert client.post('/api/uploads/commit',json={'token':'../x','company_name':'测试公司','report_title':'年度报告'}).status_code==404
    r=preview(client).json();meta=upload_preview.folder()/(r['token']+'.json')
    info=json.loads(meta.read_text(encoding='utf-8'));info['created']=time.time()-upload_preview.TTL-1;meta.write_text(json.dumps(info),encoding='utf-8')
    assert client.post('/api/uploads/commit',json={'token':r['token'],'company_name':'测试公司','report_title':'年度报告'}).status_code==410


def test_background_commit_exposes_progress_and_result(client,monkeypatch):
    info=preview(client).json();reached=Event();release=Event()
    def fake_save(file,company_name,report_title,*args,**kwargs):
        progress=kwargs.get('progress') or args[-1]
        progress(51,'精确比较 2 / 10 份原文');reached.set();assert release.wait(2)
        return {'document_id':'doc-a','company_name':company_name,'report_title':report_title}
    monkeypatch.setattr(uploads,'save_report',fake_save)
    started=client.post('/api/uploads/commit/start',json={'token':info['token'],'company_name':'甲公司','report_title':'甲报告'})
    assert started.status_code==202 and reached.wait(2)
    jid=started.json()['job_id'];running=client.get('/api/uploads/jobs/'+jid).json()
    assert running['status']=='running' and running['percent']==51 and '2 / 10' in running['stage']
    release.set()
    for _ in range(30):
        job=client.get('/api/uploads/jobs/'+jid).json()
        if job['status']=='complete':break
        time.sleep(.03)
    assert job['percent']==100 and job['result']['document_id']=='doc-a'
    assert not (upload_preview.folder()/(info['token']+'.docx')).exists()


def test_failed_background_commit_keeps_staged_file(client,monkeypatch):
    info=preview(client).json()
    def reject(*args,**kwargs):
        from fastapi import HTTPException
        (kwargs.get('progress') or args[-1])(38,'查找疑似重复报告')
        raise HTTPException(409,{'code':'exact_duplicate','message':'重复'})
    monkeypatch.setattr(uploads,'save_report',reject)
    jid=client.post('/api/uploads/commit/start',json={'token':info['token'],'company_name':'甲公司','report_title':'甲报告'}).json()['job_id']
    for _ in range(30):
        job=client.get('/api/uploads/jobs/'+jid).json()
        if job['status']=='failed':break
        time.sleep(.03)
    assert job['percent']==100 and job['error']['code']=='exact_duplicate'
    assert upload_preview.resolve(info['token'])[0].exists()
