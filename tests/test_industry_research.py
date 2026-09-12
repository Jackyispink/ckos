import asyncio

import httpx
import pytest
from pydantic import ValidationError

from app.industry_research.planner import CHAPTERS, build_plan
from app.industry_research.schemas import EvidenceCreate, ResearchBrief
from app.industry_research.search.base import SearchProvider, SearchRequest, SearchResult
from app.industry_research.search.exa import ExaSearchProvider
from app.industry_research.search.service import canonical_url, search_chapter
from app.industry_research.writer import _EVIDENCE_BUDGET, _fit_evidence


def test_regeneration_cancel_failure_and_success_preserve_original_until_commit(monkeypatch):
    from app.industry_research import database, writer
    from app.industry_research.planner import build_plan
    from app.industry_research.schemas import ResearchBrief
    database.init()
    brief = ResearchBrief(topic='重写及取消验证')
    pid = database.create_project(brief, build_plan(brief))
    for number in range(1, 11):
        database.update_chapter(pid, number, content=f'原稿第{number}章。')
        database.add_evidence(pid, EvidenceCreate(
            chapter_no=number, title=f'第{number}章测试证据', excerpt='用于验证重新生成事务边界。'))
    with database.connect() as db:
        db.execute("""UPDATE industry_evidence SET metadata=jsonb_set(metadata,'{review}',
          '{\"decision\":\"approved\",\"reason\":\"测试资料已复核\"}'::jsonb)
          WHERE project_id=%s""", (pid,))
        db.execute("INSERT INTO industry_chapters(project_id,chapter_no,title,questions,queries,content,status) VALUES (%s,0,'摘要','[]','[]','原稿摘要','complete')", (pid,))
    writer._update(pid, 'complete', '完成', progress=100, completed=10)
    monkeypatch.setattr(writer.llm, 'configured', lambda: True)
    monkeypatch.setattr(writer._pool, 'submit', lambda *args: None)
    for outcome in ('cancel', 'failure', 'success'):
        calls = []
        async def complete(messages):
            calls.append(messages)
            assert database.chapter(pid, 1)['content'] == '原稿第1章。'
            assert database.chapter(pid, 0)['content'] == '原稿摘要'
            if len(calls) == 3:
                if outcome == 'cancel':
                    assert writer.cancel(pid)
                elif outcome == 'failure':
                    raise writer.llm.ModelError('模拟重新生成失败')
            return '新稿正文，依据证据描述当前行业的主要变化。'
        monkeypatch.setattr(writer.llm, 'complete', complete)
        assert writer.start(pid, mode='regenerate')
        writer._run(pid, mode='regenerate')
        assert database.project(pid)['status'] == 'complete'
        assert not writer.is_running(pid)
        if outcome == 'success':
            assert len(calls) == 11
            versions = database.list_versions(pid)
            assert len(versions) == 1
            snapshot = database.get_version(pid, versions[0]['id'])['snapshot']
            assert snapshot['project']['chapters'][0]['content'] == '原稿摘要'
            assert database.chapter(pid, 1)['content'].startswith('新稿')
            assert database.chapter(pid, 0)['content'].startswith('新稿')
        else:
            assert len(calls) == 3
            assert database.chapter(pid, 1)['content'] == '原稿第1章。'
            assert database.chapter(pid, 0)['content'] == '原稿摘要'


def test_resume_keeps_completed_chapters_and_retries_summary_only(monkeypatch):
    from app.industry_research import database, writer
    from app.industry_research.planner import build_plan
    from app.industry_research.schemas import ResearchBrief
    database.init()
    brief = ResearchBrief(topic='继续生成验证')
    pid = database.create_project(brief, build_plan(brief))
    for number in range(1, 6):
        database.update_chapter(pid, number, content=f'已人工确认的第{number}章正文。')
    for number in range(1, 11):
        database.add_evidence(pid, EvidenceCreate(
            chapter_no=number, title=f'第{number}章测试证据', excerpt='用于验证继续生成与摘要重试。'))
    with database.connect() as db:
        db.execute("""UPDATE industry_evidence SET metadata=jsonb_set(metadata,'{review}',
          '{\"decision\":\"approved\",\"reason\":\"测试资料已复核\"}'::jsonb)
          WHERE project_id=%s""", (pid,))
        db.execute("INSERT INTO industry_chapters(project_id,chapter_no,title,questions,queries,content,status) VALUES (%s,0,'摘要','[]','[]','旧版摘要','complete')", (pid,))
    calls = []
    async def complete(messages):
        calls.append(messages)
        if len(calls) == 6:
            raise writer.llm.ModelError('摘要测试失败')
        return '这是依据现有证据撰写的研究正文，仍需持续核实行业变化。'
    monkeypatch.setattr(writer.llm, 'configured', lambda: True)
    monkeypatch.setattr(writer.llm, 'complete', complete)
    writer._run(pid)
    assert len(calls) == 6  # Chapters 6-10, then summary.
    assert database.project(pid)['status'] == 'failed'
    assert database.chapter(pid, 1)['content'] == '已人工确认的第1章正文。'
    writer._run(pid)
    assert len(calls) == 7  # Retry only summary; do not repeat the ten chapters.
    assert database.project(pid)['status'] == 'complete'
    writer._run(pid)
    assert len(calls) == 7  # Completed report is a no-op.


def test_orphan_running_report_can_resume_and_active_job_is_not_duplicated(monkeypatch):
    from app.industry_research import database, writer
    from app.industry_research.router import get_project, generate
    from app.industry_research.planner import build_plan
    from app.industry_research.schemas import ResearchBrief
    from fastapi import HTTPException
    import pytest
    database.init()
    brief = ResearchBrief(topic='中断验证')
    pid = database.create_project(brief, build_plan(brief))
    writer._update(pid, 'running', '服务退出前生成中')
    assert get_project(pid)['status'] == 'interrupted'
    monkeypatch.setattr(writer._pool, 'submit', lambda *args: None)
    try:
        assert generate(pid)['started']
        assert get_project(pid)['generation_active']
        with pytest.raises(HTTPException) as exc:
            generate(pid)
        assert exc.value.status_code == 409
    finally:
        writer._running.discard(pid)
from app.industry_research import database as industry_database
from app.industry_research.sources.akshare_provider import AkshareProvider
from app.industry_research.sources.base import DataRequirement, StructuredEvidence
from app.industry_research.sources.router import collect_chapter
from app.industry_research import collection_jobs
from app.industry_research.exporter import _clean_markdown, build_docx
from app.industry_research.content_quality import quality_issues, sanitize_generated_content
from app.industry_research.charts import discover, render


def test_plan_has_complete_ordered_report_structure():
    brief = ResearchBrief(topic='人形机器人', geography='中国', key_companies='优必选、宇树科技')
    plan = build_plan(brief)

    assert len(plan) == len(CHAPTERS) == 10
    assert [chapter['chapter_no'] for chapter in plan] == list(range(1, 11))
    assert all(chapter['questions'] and chapter['queries'] for chapter in plan)
    assert all('中国人形机器人' in query for query in plan[0]['queries'])
    assert any('优必选' in query for query in plan[7]['queries'])


def test_research_period_must_be_consistent():
    with pytest.raises(ValidationError):
        ResearchBrief(topic='低空经济', history_start=2026, history_end=2025)
    with pytest.raises(ValidationError):
        ResearchBrief(topic='低空经济', history_end=2025, forecast_end=2025)


def test_exa_adapter_maps_vendor_response_to_common_result():
    def handler(request):
        assert request.headers['x-api-key'] == 'secret'
        payload = __import__('json').loads(request.content)
        assert payload['query'] == '中国人形机器人市场规模'
        return httpx.Response(200, json={'results': [{
            'id': 'exa-1', 'title': '市场报告', 'url': 'https://example.com/report?utm_source=test',
            'publishedDate': '2026-01-02T00:00:00Z', 'author': '行业协会',
            'summary': '2025年市场规模及其统计口径。', 'highlights': ['关键数据'], 'highlightScores': [0.82],
        }]})

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            provider = ExaSearchProvider(api_key='secret', client=client)
            return await provider.search(SearchRequest(query='中国人形机器人市场规模'))

    results = asyncio.run(run())
    assert results[0].provider == 'exa'
    assert results[0].publisher == '行业协会'
    assert results[0].score == .82
    assert results[0].text == '2025年市场规模及其统计口径。'


def test_common_search_service_deduplicates_urls(monkeypatch):
    class FakeProvider(SearchProvider):
        name = 'fake'

        async def search(self, request):
            return [SearchResult('资料', 'https://www.example.com/a?utm_source=x', '证据', provider='fake')]

    monkeypatch.setattr('app.industry_research.search.service.database.chapter',
                        lambda project_id, chapter_no: {'queries': ['查询一', '查询二']})
    monkeypatch.setattr('app.industry_research.search.service.database.all_evidence', lambda project_id: [])
    saved = []
    monkeypatch.setattr('app.industry_research.search.service.database.add_search_results',
                        lambda project_id, chapter_no, rows: saved.extend(rows) or len(rows))
    result = asyncio.run(search_chapter('p1', 4, max_queries=2, provider=FakeProvider()))
    assert result['found'] == result['saved'] == 1
    assert canonical_url('https://www.example.com/a?utm_source=x') == 'https://example.com/a'


def test_chapter_evidence_is_ranked_and_bounded():
    evidence = [{'title': f'来源{i}', 'excerpt': '资料' * 1000, 'score': i / 100} for i in range(30)]
    chosen = _fit_evidence(evidence)
    assert chosen[0]['score'] == .29
    assert sum(len(x['excerpt']) + len(x['title']) + 80 for x in chosen) <= _EVIDENCE_BUDGET
    assert all(len(x['excerpt']) <= 1200 for x in chosen)


def test_akshare_adapter_normalizes_dataframe(monkeypatch):
    import pandas as pd

    monkeypatch.setattr(AkshareProvider, '_call', staticmethod(lambda spec, params: pd.DataFrame([
        {'季度': '2025年第1季度', '国内生产总值-绝对值': 100},
        {'季度': '2025年第2季度', '国内生产总值-绝对值': 210},
    ])))
    evidence = asyncio.run(AkshareProvider().fetch(DataRequirement('china_gdp', '经济增长', 3)))
    assert evidence.provider == 'akshare'
    assert evidence.dataset == 'china_gdp'
    assert '2025年第2季度' in evidence.text
    assert evidence.metadata['rows_returned'] == 2


def test_source_router_uses_structured_data_before_reducing_web_budget(monkeypatch):
    class FakeStructured:
        name = 'akshare'

        async def fetch(self, requirement):
            return StructuredEvidence(requirement.label, '结构化证据', 'https://example.com', 'akshare', requirement.key)

    class FakeSearch(SearchProvider):
        name = 'fake-search'

        async def search(self, request):
            return []

    monkeypatch.setattr('app.industry_research.sources.router.database.project',
                        lambda project_id: {'brief': {'geography': '中国'}})
    monkeypatch.setattr('app.industry_research.sources.router.database.chapter',
                        lambda project_id, chapter_no: {'queries': ['政策', '技术', '经济']})
    monkeypatch.setattr('app.industry_research.sources.router.database.evidence_urls', lambda project_id: [])
    monkeypatch.setattr('app.industry_research.search.service.database.all_evidence', lambda project_id: [])
    monkeypatch.setattr('app.industry_research.sources.router.database.get_source_cache', lambda *args: None)
    monkeypatch.setattr('app.industry_research.sources.router.database.put_source_cache', lambda *args: None)
    monkeypatch.setattr('app.industry_research.sources.router.database.add_structured_evidence', lambda *args: 1)
    monkeypatch.setattr('app.industry_research.sources.router.database.add_search_results', lambda *args: 0)
    result = asyncio.run(collect_chapter('p1', 3, max_web_queries=3,
                                         structured_provider=FakeStructured(), search_provider=FakeSearch()))
    assert result['structured_saved'] == 3
    assert result['web_query_budget'] == 2
    assert result['web']['queries_run'] == 2


def test_project_and_all_chapters_are_persisted():
    industry_database.init()
    brief = ResearchBrief(topic='AI短剧', geography='中国')
    project_id = industry_database.create_project(brief, build_plan(brief))
    project = industry_database.project(project_id)
    assert project['title'] == '中国AI短剧行业研究'
    assert len(project['chapters']) == 10


def test_identical_empty_research_plan_is_reused():
    industry_database.init()
    brief = ResearchBrief(topic='AI短剧', geography='中国')
    first = industry_database.create_project(brief, build_plan(brief))
    second = industry_database.create_project(brief, build_plan(brief))
    assert second == first


def test_collection_job_exposes_progress(monkeypatch):
    async def fake_collect(project_id, chapter_no, max_web_queries, results_per_query, progress):
        progress('web', 50, '网页搜索 1/2：测试词', 1, 2, 4)
        return {'structured_saved': 1, 'web': {'saved': 4}}

    monkeypatch.setattr(collection_jobs, 'collect_chapter', fake_collect)
    started = collection_jobs.start('p1', 1, 2, 8)
    import time
    for _ in range(30):
        job = collection_jobs.get(started['job_id'])
        if job['status'] == 'complete':
            break
        time.sleep(.01)
    assert job['percent'] == 100
    assert job['saved'] == 5
    assert job['message'] == '资料采集完成'


@pytest.mark.parametrize('result', [
    {'structured_saved': 1, 'web': {'saved': 4, 'queries_run': 2},
     'web_error': '对端未返回HTTP响应便断开'},
    {'structured_saved': 0, 'web': {'saved': 3, 'queries_run': 1, 'partial': True}},
])
def test_collection_job_marks_partial_web_search_as_warning(monkeypatch, result):
    async def fake_collect(*args, **kwargs):
        return result

    monkeypatch.setattr(collection_jobs, 'collect_chapter', fake_collect)
    started = collection_jobs.start('p1', 2, 3, 8)
    import time
    for _ in range(30):
        job = collection_jobs.get(started['job_id'])
        if job['status'] == 'complete_with_warning':
            break
        time.sleep(.01)

    assert job['status'] == 'complete_with_warning'
    assert job['stage'] == 'complete_with_warning'
    assert job['saved'] == result['structured_saved'] + result['web']['saved']
    assert job['result'] == result
    assert '网页搜索中断' in job['message']
    assert f"已保留本次新增 {job['saved']} 条" in job['message']
    assert job['warning'] == result.get('web_error')


def test_completed_report_exports_as_styled_docx(tmp_path, monkeypatch):
    from docx import Document
    project = {'id': 'report1', 'title': '中国AI短剧行业研究', 'brief': {
        'geography': '中国', 'history_start': 2023, 'history_end': 2026, 'forecast_end': 2030,
        'purpose': '行业战略判断', 'focus': '市场规模', 'included_segments': '', 'excluded_segments': '', 'key_companies': ''},
        'chapters': [{'chapter_no': 0, 'title': '摘要', 'content': '## 核心观点\n- 行业处于成长期。'}] +
                    [{'chapter_no': i, 'title': f'章节{i}',
                      'content': f'第{i}章正文。' + ('关键事实来自资料。[1]' if i == 1 else '')}
                     for i in range(1, 11)]}
    monkeypatch.setattr('app.industry_research.exporter.EXPORT_DIR', tmp_path)
    path = build_docx(project, [{'chapter_no': 1, 'title': '来源资料', 'publisher': '机构',
                                 'published_at': '2026', 'url': 'https://example.com',
                                 'excerpt': '支持第1章关键事实的原文。'}])
    doc = Document(path)
    assert doc.core_properties.title == project['title']
    assert any(p.text == '摘要：核心结论' for p in doc.paragraphs)
    assert any(p.text == '参考资料' for p in doc.paragraphs)


def test_exporter_removes_duplicate_chapter_heading_and_all_markdown_hashes(tmp_path, monkeypatch):
    from docx import Document
    project = {'id': 'clean-report', 'title': '行业研究', 'brief': {
        'geography': '中国', 'history_start': 2023, 'history_end': 2026, 'forecast_end': 2030,
        'purpose': '', 'focus': '', 'included_segments': '', 'excluded_segments': '', 'key_companies': ''},
        'chapters': [{'chapter_no': 0, 'title': '摘要', 'content': '#### 核心观点\n正文'}] +
                    [{'chapter_no': i, 'title': f'章节{i}',
                      'content': f'## 第{i}章 章节{i}\n#### 结论\n正文'} for i in range(1, 11)]}
    monkeypatch.setattr('app.industry_research.exporter.EXPORT_DIR', tmp_path)
    path = build_docx(project, [])
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    assert not any('#' in text for text in texts)
    assert sum(p.text == '第1章  章节1' and p.style.name == 'Heading 1' for p in doc.paragraphs) == 1
    assert any(p.text.endswith(' 结论') and p.style.name == 'Heading 2' for p in doc.paragraphs)
    assert _clean_markdown('## 第1章 章节1\n正文', '章节1').strip() == '正文'


def test_internal_prompt_leakage_is_removed_without_losing_real_evidence_gap():
    content = '''## 结论
证据不足，当前无法判断市场份额。公开资料尚未给出统一统计口径，各平台披露的数据也无法直接横向比较。因此，本报告将市场份额列为待验证事项，现阶段仅讨论参与者类型、产品能力和商业模式，不对企业排名作无依据推断。

## **重要事实用[编号]引用；没有证据的数字不得编造。**
## **证据不足时明确写出“证据不足”和需要补充的数据，不要用常识填充。**
## **使用中文Markdown，控制在1200字以内。**'''
    cleaned = sanitize_generated_content(content)
    assert '当前无法判断市场份额' in cleaned
    assert '不得编造' not in cleaned
    assert '1200字' not in cleaned
    assert quality_issues(content) == ['模型复述了内部写作指令']
    assert quality_issues(cleaned) == []


def test_chart_candidates_are_extracted_and_rendered_without_llm(tmp_path, monkeypatch):
    project = {'chapters': [{'chapter_no': 4, 'title': '市场规模', 'content':
        '2022年市场规模为100亿元。\n2023年市场规模达到150亿元。\n2024年市场规模为210亿元。'}]}
    candidates = discover(project)
    assert candidates[0]['labels'] == ['2022', '2023', '2024']
    assert candidates[0]['values'] == [100.0, 150.0, 210.0]
    assert candidates[0]['chart_types'] == ['line', 'bar', 'column', 'area', 'scatter', 'lollipop']
    monkeypatch.setattr('app.industry_research.charts.CHART_DIR', tmp_path)
    candidates[0]['selected_type'] = 'bar'
    assert render('p1', candidates[0]).is_file()


def test_report_crud_and_chart_choices_are_persisted():
    industry_database.init(); brief = ResearchBrief(topic='储能', geography='中国')
    project_id = industry_database.create_project(brief, build_plan(brief))
    assert industry_database.update_project(project_id, '中国储能行业深度研究') == 1
    assert industry_database.update_chapter(project_id, 1, content='2024年市场规模为100亿元。') == 1
    choice = type('Choice', (), {'id': 'candidate123', 'chart_type': 'bar', 'title': '市场规模', 'selected': True})()
    assert industry_database.save_chart_selections(project_id, [choice])
    assert industry_database.chart_selections(project_id)['candidate123']['chart_type'] == 'bar'
    assert industry_database.delete_project(project_id) == 1
    assert industry_database.project(project_id) is None
