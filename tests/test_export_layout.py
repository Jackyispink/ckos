from docx import Document
from docx.oxml.ns import qn
from app.industry_research.exporter import _markdown, _reference_material, _sources, _styles


def test_lists_restart_per_group_and_keep_internal_sequence():
    doc = Document(); _styles(doc)
    _markdown(doc, '1. 第一项\n2. 第二项\n## 下一小节\n1. 新组第一项\n2. 新组第二项')
    items = [p for p in doc.paragraphs if p.style.name == 'List Number']
    ids = [p._p.pPr.numPr.numId.val for p in items]
    assert ids[0] == ids[1] and ids[2] == ids[3] and ids[0] != ids[2]
    assert items[0].runs[0]._r.rPr.rFonts.get(qn('w:eastAsia')) == 'SimSun'


def test_reference_merge_preserves_analysis_and_unique_entries():
    evidence = [dict(chapter_no=n, title='共同来源', url='https://example.com/report', publisher='机构', published_at='2026-01-01T00:00:00') for n in (1, 2)]
    chapters = [dict(chapter_no=n, content='## 数据来源与统计口径\n必须保留的统计解释[1]。\n## 参考文献\n[1] 共同来源｜机构\n[2] 独有线下资料\n## 风险\n必须保留的风险解释。') for n in (1, 2)]
    cleaned, refs = _reference_material(chapters, evidence)
    assert '统计解释' in cleaned[0]['content'] and '风险解释' in cleaned[0]['content']
    assert '共同来源' not in cleaned[0]['content']
    doc = Document(); _sources(doc, evidence, refs)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert text.count('共同来源') == 1
    assert text.count('独有线下资料') == 1
    assert '1-1' in text and '2-1' in text
    assert 'https://' not in text
    assert any(r.target_ref == 'https://example.com/report' for r in doc.part.rels.values())


def test_export_smoke_without_database(tmp_path, monkeypatch):
    from app.industry_research import exporter
    monkeypatch.setattr(exporter, 'EXPORT_DIR', tmp_path)
    project = dict(id='layout-test', title='版式验证', brief={}, chapters=[dict(chapter_no=1, title='行业分析', content='## 行业分析\n1. 首项\n2. 次项\n## 参考文献\n[1] 来源甲\n## 风险\n保留风险[1]。')])
    path = exporter.build_docx(project, [])
    doc = Document(path)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert '保留风险' in text
    assert text.count('来源甲') == 1
    assert '参考文献' not in text
    assert '参考资料' in text


def test_body_citations_receive_chapter_prefix():
    doc=Document();_styles(doc)
    _markdown(doc,'关键数字为100亿元[1]。',chapter_no=4)
    assert doc.paragraphs[0].text == '关键数字为100亿元[4-1]。'
    assert doc.paragraphs[0].runs[1].font.superscript


def test_chart_anchors_follow_source_and_number_in_reading_order(monkeypatch):
    from app.industry_research import exporter
    doc=Document()
    _markdown(doc,'2023年市场规模100亿元。\n2024年市场规模120亿元。\n后续分析文字。')
    blocks=[x for x in doc._element.body if x.tag!=qn('w:sectPr')]
    def marker(doc,specs,start=1):
        doc.add_paragraph(f'图{start}：'+specs[0]['title'])
    monkeypatch.setattr(exporter,'_chapter_dashboard',marker)
    exporter._place_chapter_charts(doc,blocks,[
        dict(title='趋势',source_excerpts=['2023年市场规模100亿元。','2024年市场规模120亿元。']),
        dict(title='首年',source_excerpts=['**2023年市场规模100亿元。**[1]']),
        dict(title='无法匹配',source_excerpts=['原文已经修改']),
        dict(title='同段第二图',source_excerpts=['2023年市场规模100亿元。'])])
    assert [p.text for p in doc.paragraphs]==['2023年市场规模100亿元。','图1：首年','图2：同段第二图','2024年市场规模120亿元。','图3：趋势','后续分析文字。','本章相关测算或程序测算结果如图 4所示。','图4：无法匹配']


def test_chart_table_and_ambiguous_excerpt(monkeypatch):
    from app.industry_research import exporter
    doc=Document();_markdown(doc,'重复的来源文字。\n重复的来源文字。\n| 年份 | 规模 |\n| --- | --- |\n| 2024 | 120亿元 |\n结束分析。')
    blocks=[x for x in doc._element.body if x.tag!=qn('w:sectPr')]
    monkeypatch.setattr(exporter,'_chapter_dashboard',lambda doc,specs,start=1:doc.add_paragraph(specs[0]['title']))
    exporter._place_chapter_charts(doc,blocks,[dict(title='表格图',source_excerpts=['| 2024 | 120亿元 |']),dict(title='歧义图',source_excerpts=['重复的来源文字。'])])
    nodes=list(doc._element.body)
    table_index=next(i for i,n in enumerate(nodes) if n.tag==qn('w:tbl'))
    assert ''.join(nodes[table_index+1].xpath('.//w:t/text()'))=='表格图'
    assert doc.paragraphs[-1].text=='歧义图'


def test_inline_figure_references_share_caption_order(monkeypatch):
    from app.industry_research import exporter
    doc=Document();_markdown(doc,'市场规模为100亿元[1]。\n后文讨论。')
    blocks=[x for x in doc._element.body if x.tag!=qn('w:sectPr')]
    monkeypatch.setattr(exporter,'_chapter_dashboard',lambda doc,specs,start=1:doc.add_paragraph(f"图 {specs[0]['chapter_no']}-{start}"))
    specs=[dict(chapter_no=4,title='缺失',source_excerpts=['缺失的来源文字']),dict(chapter_no=4,title='规模',source_excerpts=['市场规模为100亿元[1]。']),dict(chapter_no=4,title='规模二',source_excerpts=['市场规模为100亿元[1]。'])]
    exporter._place_chapter_charts(doc,blocks,specs)
    assert doc.paragraphs[0].text=='市场规模为100亿元[1]。相关数据如图 4-1、图 4-2所示。'
    assert [p.text for p in doc.paragraphs[1:]]==['图 4-1','图 4-2','后文讨论。','本章相关测算或程序测算结果如图 4-3所示。','图 4-3']
    assert doc.paragraphs[0].runs[1].font.superscript


def test_table_figure_reference_precedes_caption(monkeypatch):
    from app.industry_research import exporter
    doc=Document();_markdown(doc,'| 年份 | 规模 |\n| --- | --- |\n| 2024 | 120亿元 |')
    blocks=[x for x in doc._element.body if x.tag!=qn('w:sectPr')]
    monkeypatch.setattr(exporter,'_chapter_dashboard',lambda doc,specs,start=1:doc.add_paragraph('图 4-1'))
    exporter._place_chapter_charts(doc,blocks,[dict(chapter_no=4,title='规模',source_excerpts=['| 2024 | 120亿元 |'])])
    nodes=list(doc._element.body)
    assert nodes[0].tag==qn('w:tbl')
    assert ''.join(nodes[1].xpath('.//w:t/text()'))=='相关数据如图 4-1所示。'
    assert ''.join(nodes[2].xpath('.//w:t/text()'))=='图 4-1'


def test_markdown_list_chart_sources_match_rendered_word(monkeypatch):
    from app.industry_research import exporter
    doc=Document();_markdown(doc,'- 2021年微短剧市场：3.68亿元\n- 2022年：101.7亿元\n后续分析。')
    blocks=[x for x in doc._element.body if x.tag!=qn('w:sectPr')]
    monkeypatch.setattr(exporter,'_chapter_dashboard',lambda doc,specs,start=1:doc.add_paragraph(f'图 7-{start}'))
    exporter._place_chapter_charts(doc,blocks,[dict(chapter_no=7,title='市场',source_excerpts=['- 2021年微短剧市场：3.68亿元','- 2022年：101.7亿元'])])
    assert doc.paragraphs[1].text.endswith('相关数据如图 7-1所示。')
    assert doc.paragraphs[2].text=='图 7-1'
    assert exporter._anchor_text('-3.5亿元')=='-3.5亿元'
