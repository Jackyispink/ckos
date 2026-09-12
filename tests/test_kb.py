from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app import kb, llm
from app.main import app


def prepare(tmp_path, monkeypatch):
    monkeypatch.setattr(kb, 'ROOT', tmp_path)
    (tmp_path / 'data').mkdir()
    doc = Document()
    doc.add_paragraph('甲公司研发管理问题是缺少研发预算和研发项目复盘。')
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = '研发投入 500 万元'
    doc.save(tmp_path / 'data' / '甲公司.docx')
    return kb.ingest()


def test_ingest_table_and_idempotency(tmp_path, monkeypatch):
    assert prepare(tmp_path, monkeypatch)['indexed'] == 1
    assert kb.ingest()['unchanged'] == 1
    rows = kb.Retriever().search('研发投入')
    assert '500' in rows[0]['text']
    assert '表 1' in rows[0]['location']
    assert kb.Retriever().search('研发投入', 'nonexistent') == []


def test_sessions_persistence_and_failure(tmp_path, monkeypatch):
    prepare(tmp_path, monkeypatch)
    monkeypatch.setattr(llm, 'configured', lambda: False)
    with TestClient(app) as client:
        sid = client.post('/api/sessions').json()['id']
        other = client.post('/api/sessions').json()['id']
        res = client.post(f'/api/sessions/{sid}/ask', json={'question': '甲公司研发管理有哪些问题？'})
        assert res.status_code == 200
        assert res.json()['sources']
        assert len(client.get(f'/api/sessions/{sid}').json()) == 2
        assert client.get(f'/api/sessions/{other}').json() == []
        calls = []
        async def fake(messages):
            calls.append(messages)
            return '甲公司研发管理如何改进' if len(calls) == 1 else '建议建立预算制度 [1]'
        monkeypatch.setattr(llm, 'configured', lambda: True)
        monkeypatch.setattr(llm, 'complete', fake)
        response = client.post(f'/api/sessions/{sid}/ask', json={'question': '那应该怎么改进？'})
        assert response.status_code == 200
        assert len(calls) == 2
        assert any('甲公司研发管理有哪些问题' in m['content'] for m in calls[0])
        assert '本轮资料' in calls[1][-1]['content']
        async def failed(messages):
            raise RuntimeError('secret')
        monkeypatch.setattr(llm, 'complete', failed)
        response = client.post(f'/api/sessions/{sid}/ask', json={'question': '进一步解释'})
        assert response.status_code == 502
        assert 'secret' not in response.text
        assert len(client.get(f'/api/sessions/{sid}').json()) == 4
    with TestClient(app) as client:
        assert len(client.get(f'/api/sessions/{sid}').json()) == 4


def test_delete_session_only_removes_its_history(tmp_path, monkeypatch):
    prepare(tmp_path, monkeypatch)
    monkeypatch.setattr(llm, 'configured', lambda: False)
    with TestClient(app) as client:
        first = client.post('/api/sessions').json()['id']
        other = client.post('/api/sessions').json()['id']
        for sid in (first, other):
            assert client.post(f'/api/sessions/{sid}/ask', json={'question': '研发管理问题'}).status_code == 200
        assert client.delete(f'/api/sessions/{first}').status_code == 200
        assert client.get(f'/api/sessions/{first}').status_code == 404
        assert client.delete(f'/api/sessions/{first}').status_code == 404
        assert client.post(f'/api/sessions/{first}/ask', json={'question': '继续'}).status_code == 404
        assert len(client.get(f'/api/sessions/{other}').json()) == 2
        with kb.connect() as db:
            assert db.execute('SELECT count(*) AS n FROM messages WHERE session_id=%s', (first,)).fetchone()['n'] == 0
        assert len(client.get('/api/status').json()['documents']) == 1


def test_company_overview_reads_beyond_top_ten(tmp_path, monkeypatch):
    prepare(tmp_path, monkeypatch)
    doc = Document()
    for i in range(20):
        doc.add_paragraph(f'模块{i}的改善建议：' + '建立目标并定期复盘。' * 80)
    doc.save(tmp_path / 'data' / '04中稀（常熟）稀土新材料有限公司管理诊断报告.docx')
    kb.ingest()
    r = kb.Retriever()
    sources, coverage = r.context('中稀(常熟)稀土新材料有限公司的诊断建议是什么')
    assert len(sources) > 10
    assert '全部已入库' in coverage
    assert all('中稀' in s['name'] for s in sources)
    assert any('模块19' in s['text'] for s in sources)
    focused, _ = r.context('中稀(常熟)稀土新材料有限公司模块19是什么')
    assert all('中稀' in s['name'] for s in focused)
