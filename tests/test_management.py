from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient
from app import database,kb,structured
from app.main import app


def test_management_lifecycle(tmp_path,monkeypatch):
    monkeypatch.setattr(database,'ROOT',tmp_path);monkeypatch.setattr(kb,'ROOT',tmp_path)
    d=Document();d.add_paragraph('测试企业');d.add_heading('人力资源管理诊断与建议',level=2);d.add_paragraph('改善建议');d.add_paragraph('原建议');buf=BytesIO();d.save(buf)
    with TestClient(app) as client:
        r=client.post('/api/documents/upload',data={'company_name':'甲甲公司','report_title':'报告一号'},files={'file':('x.docx',buf.getvalue())});assert r.status_code==201,r.text
        did=r.json()['document_id'];base='/api/manage/reports/'+did
        assert client.patch(base+'/names',json={'company_name':'乙乙公司','report_title':'修改后报告'}).status_code==200
        assert client.get('/api/status').json()['documents'][0]['name']=='修改后报告'
        assert not structured.resolve('甲甲公司建议')[0]
        items=client.get(base+'/contents').json();bid=next(b['id'] for b in items if b['text']=='原建议')
        body={'module':'人力资源','section':'建议','text':'修改后独有建议'}
        assert client.put('/api/manage/contents/'+bid,json=body).status_code==200
        assert '修改后独有建议' in structured.resolve('乙乙公司建议')[0][0]['text']
        added=client.post(base+'/contents',json=dict(body,text='新增建议')).json()['id']
        assert client.patch('/api/manage/contents/'+added+'/state',json={'deleted':True}).status_code==200
        assert '新增建议' not in structured.resolve('乙乙公司建议')[0][0]['text']
        client.patch('/api/manage/contents/'+added+'/state',json={'deleted':False})
        assert '新增建议' in structured.resolve('乙乙公司建议')[0][0]['text']
        assert structured.ingest(tmp_path)['unchanged']==1
        assert '修改后独有建议' in structured.resolve('乙乙公司建议')[0][0]['text']
        client.patch(base+'/state',json={'deleted':True})
        assert not structured.resolve('乙乙公司建议')[0]
        assert client.get('/api/status').json()['documents']==[]
        assert client.get('/api/preprocessing').json()==[]
        duplicate=client.post('/api/documents/upload',data={'company_name':'乙乙公司','report_title':'另一个名字'},files={'file':('renamed.docx',buf.getvalue())}).json()
        assert duplicate['detail']['code']=='exact_duplicate' and duplicate['detail']['matches'][0]['deleted']
        client.patch(base+'/state',json={'deleted':False})
        assert structured.resolve('乙乙公司建议')[0]
        with database.connect() as db:
            assert db.execute('SELECT count(*) n FROM content_audit').fetchone()['n']>=6
