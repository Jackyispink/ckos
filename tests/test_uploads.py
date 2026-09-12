from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app import database, kb, structured
from app.main import app


def document_bytes():
    d=Document();d.add_paragraph('错误识别有限公司');d.add_heading('人力资源管理诊断与建议',level=2)
    d.add_paragraph('存在现状');d.add_paragraph('人才梯队不足。');d.add_paragraph('改善建议');d.add_paragraph('制定培养计划。')
    buf=BytesIO();d.save(buf);return buf.getvalue()


def company_document(company,marker):
    d=Document();d.add_paragraph(company);d.add_heading('人力资源管理诊断与建议',level=2)
    d.add_paragraph('存在现状');d.add_paragraph(marker+'人才问题。');d.add_paragraph('改善建议');d.add_paragraph(marker+'培养方案。')
    buf=BytesIO();d.save(buf);return buf.getvalue()


def test_upload_user_names_and_persistence(tmp_path,monkeypatch):
    monkeypatch.setattr(kb,'ROOT',tmp_path);monkeypatch.setattr(database,'ROOT',tmp_path)
    with TestClient(app) as client:
        body={'company_name':'用户指定有限公司','report_title':'用户定义报告'}
        payload=document_bytes()
        result=client.post('/api/documents/upload',data=body,files={'file':('../../同名.docx',payload)})
        assert result.status_code==201,result.text
        did=result.json()['document_id']
        assert client.get('/api/status').json()['documents'][0]['name']=='用户定义报告'
        sources,_=structured.resolve('用户指定有限公司人才问题')
        assert sources and sources[0]['document_id']==did
        assert not structured.resolve('错误识别有限公司人才问题')[0]
        assert client.get('/api/documents/'+did).status_code==200
        assert structured.ingest(tmp_path)['unchanged']==1
        result2=client.post('/api/documents/upload',data=body,files={'file':('同名.docx',payload)})
        assert result2.status_code==409 and result2.json()['detail']['code']=='name_conflict'
        assert len(list((tmp_path/'data').glob('*.docx')))==1
        assert not (tmp_path/'同名.docx').exists()


def test_invalid_upload_not_published(tmp_path,monkeypatch):
    monkeypatch.setattr(kb,'ROOT',tmp_path);monkeypatch.setattr(database,'ROOT',tmp_path)
    with TestClient(app) as client:
        body={'company_name':'测试企业','report_title':'报告测试'}
        for name,data in [('fake.docx',b'not zip'),('bad.pdf',b'pdf')]:
            assert client.post('/api/documents/upload',data=body,files={'file':(name,data)}).status_code==422
        assert client.get('/api/status').json()['documents']==[]
        assert not list((tmp_path/'data').glob('*.docx'))


def test_uploaded_reports_are_isolated_at_every_storage_and_retrieval_layer(tmp_path,monkeypatch):
    monkeypatch.setattr(kb,'ROOT',tmp_path);monkeypatch.setattr(database,'ROOT',tmp_path)
    with TestClient(app) as client:
        ids=[]
        for company,title,marker in [('甲方制造有限公司','甲方诊断','甲方独有证据'),('乙方材料有限公司','乙方诊断','乙方独有证据')]:
            response=client.post('/api/documents/upload',data={'company_name':company,'report_title':title},files={'file':(title+'.docx',company_document(company,marker))})
            assert response.status_code==201,response.text
            ids.append(response.json()['document_id'])
        assert ids[0]!=ids[1]
        with database.connect() as db:
            chunks={did:'\n'.join(r['text'] for r in db.execute('SELECT text FROM chunks WHERE document_id=%s',(did,))) for did in ids}
            reports=[db.execute('SELECT id,company_id FROM kb_reports WHERE document_id=%s AND active',(did,)).fetchone() for did in ids]
            blocks={r['id']:'\n'.join(x['text'] for x in db.execute('SELECT text FROM kb_blocks WHERE report_id=%s',(r['id'],))) for r in reports}
            indexed=[db.execute('SELECT document_id FROM dedup_documents WHERE document_id=%s',(did,)).fetchone() for did in ids]
        assert '甲方独有证据' in chunks[ids[0]] and '乙方独有证据' not in chunks[ids[0]]
        assert '乙方独有证据' in chunks[ids[1]] and '甲方独有证据' not in chunks[ids[1]]
        assert reports[0]['company_id']!=reports[1]['company_id'] and all(indexed)
        assert '甲方独有证据' in blocks[reports[0]['id']] and '乙方独有证据' not in blocks[reports[0]['id']]
        sources,_=structured.resolve('甲方制造有限公司的人才建议',document_id=ids[0])
        assert sources and {s['document_id'] for s in sources}=={ids[0]}
        assert '乙方独有证据' not in '\n'.join(s['text'] for s in sources)
        mismatch,message=structured.resolve('乙方材料有限公司的人才建议',document_id=ids[0])
        assert mismatch==[] and '不一致' in message
