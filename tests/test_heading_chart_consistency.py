from docx import Document

from app.industry_research import charts, exporter


def test_numbered_and_unnumbered_chapters_have_same_relative_hierarchy():
    doc = Document()
    exporter._markdown(doc, '## 第1章 定义\n### 1.1 行业边界\n正文。\n### 1.2 研究范围\n#### 纳入范围\n正文。\n#### 排除范围\n正文。\n### 1.3 结论\n结论正文。', '定义', 1)
    exporter._markdown(doc, '## 第2章 历程\n### 关键事件\n#### 技术突破\n正文。\n#### 政策变化\n正文。\n### 生命周期\n正文。\n### 结论\n结论正文。', '历程', 2)
    headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith('Heading')]
    assert headings == [
        ('Heading 2', '1.1 行业边界'), ('Heading 2', '1.2 研究范围'),
        ('Heading 3', '1.2.1 纳入范围'), ('Heading 3', '1.2.2 排除范围'),
        ('Heading 2', '1.3 结论'), ('Heading 2', '2.1 关键事件'),
        ('Heading 3', '2.1.1 技术突破'), ('Heading 3', '2.1.2 政策变化'),
        ('Heading 2', '2.2 生命周期'), ('Heading 2', '2.3 结论')]


def test_heading_numbers_preserve_years_and_ordinary_lists():
    doc = Document()
    exporter._markdown(doc, '## 2026年展望\n1. 第一条\n2. 第二条\n## 3D技术\n正文。', chapter_no=4)
    assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 2'] == ['4.1 2026年展望', '4.2 3D技术']
    assert [p.text for p in doc.paragraphs if p.style.name == 'List Number'] == ['第一条', '第二条']


def series(content):
    return [s for s in charts.discover({'chapters': [dict(chapter_no=4, title='市场', content=content)]}) if s['kind'] == 'time_series']


def test_actual_report_three_years_shorthand_is_one_series():
    result = series('#### 增速预测\n- **2024年**：市场规模突破500亿元，增速约34.9%。\n- **2025年**：规模约677.9亿元，增速约34.4%。\n- **2026年**：预计793.3亿元，增速约17.0%。[10]')
    assert len(result) == 1
    assert result[0]['labels'] == ['2024', '2025', '2026']
    assert result[0]['values'] == [500, 677.9, 793.3]
    assert len(result[0]['source_excerpts']) == 3
    assert '突破500' in result[0]['source_excerpts'][0]


def test_manufacturing_area_series_keeps_all_three_years():
    result = series('2023年目标产品需求100万㎡。\n2024年需求120万㎡。\n2025年需求150万㎡。')
    assert len(result) == 1
    assert result[0]['unit'] == '万㎡'
    assert result[0]['labels'] == ['2023', '2024', '2025']
    assert result[0]['values'] == [100, 120, 150]


def test_free_text_series_keeps_complete_thousands_separators():
    result = series(
        '2023年市场规模1,200亿元。\n'
        '2024年规模1，500亿元。'
    )

    assert len(result) == 1
    assert result[0]['labels'] == ['2023', '2024']
    assert result[0]['values'] == [1200, 1500]


def test_ambiguous_respective_year_values_are_not_chart_candidates():
    found = charts.discover({'chapters': [{
        'chapter_no': 4,
        'title': '市场',
        'content': '2022年和2023年市场规模分别为100亿元和120亿元。',
    }]})

    assert found == []


def test_shorthand_never_merges_named_markets_units_or_sections():
    result = series('### 微短剧\n2024年微短剧市场规模500亿元。\n2025年规模677亿元。\n2024年AI漫剧市场规模50亿元。\n2025年规模100亿元。\n### 海外\n2026年规模100亿元。\n2027年规模200亿元。')
    assert [s['values'] for s in result] == [[500, 677], [50, 100], [100, 200]]
    assert series('2024年市场规模500亿元。\n2025年规模600万亿元。') == []


def test_export_numbers_all_chapters_without_changing_source(tmp_path, monkeypatch):
    monkeypatch.setattr(exporter, 'EXPORT_DIR', tmp_path)
    chapters = [dict(chapter_no=n, title=f'研究{n}', content='### 分析\n正文。\n#### 依据\n证据。') for n in range(1, 11)]
    doc = Document(exporter.build_docx(dict(id='headings', title='标题验证', brief={}, chapters=chapters), []))
    assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 2'] == [f'{n}.1 分析' for n in range(1, 11)]
    assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 3'] == [f'{n}.1.1 依据' for n in range(1, 11)]
    assert all(c['content'].startswith('### 分析') for c in chapters)
