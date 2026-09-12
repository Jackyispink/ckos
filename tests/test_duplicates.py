from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import database, duplicates, kb
from app.main import app


def word(change='', bold=False):
    doc = Document()
    doc.add_heading('人力资源管理诊断与建议', level=2)
    doc.add_paragraph('改善建议')
    for i in range(50):
        doc.add_paragraph(f'第{i}项独立改进建议：为部门{i}建立绩效评价表，按照第{i+10}号制度开展季度培训与能力考核。').runs[0].bold = bold
    doc.add_paragraph('设备维修预算为12.5万元。' + change)
    buffer = BytesIO(); doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(kb, 'ROOT', tmp_path)
    monkeypatch.setattr(database, 'ROOT', tmp_path)
    with TestClient(app) as c:
        yield c


def upload(client, payload, title='年度报告', company='测试（甲）公司', **fields):
    return client.post('/api/documents/upload', data={'company_name': company, 'report_title': title, **fields}, files={'file': ('report.docx', payload)})


def test_name_checked_before_extract_and_chunk(client, monkeypatch, tmp_path):
    assert upload(client, word()).status_code == 201
    def forbidden(*args, **kwargs):
        raise AssertionError('must reject before extraction or chunking')
    monkeypatch.setattr(duplicates, 'extract', forbidden)
    monkeypatch.setattr(kb, 'ingest', forbidden)
    response = upload(client, b'invalid', company=' 测试(甲)公司 ', title='年度 报告.docx')
    assert response.status_code == 409 and response.json()['detail']['code'] == 'name_conflict'
    assert len(list((tmp_path/'data').glob('*.docx'))) == 1
    assert client.post('/api/uploads/check-name', json={'company_name':'另一家公司','report_title':'年度报告'}).status_code == 200


def test_format_only_duplicate_rejected_before_chunking(client, monkeypatch, tmp_path):
    assert upload(client, word()).status_code == 201
    monkeypatch.setattr(kb, 'ingest', lambda *args: pytest.fail('duplicate must not be chunked'))
    response = upload(client, word(bold=True), title='改名后报告', similarity_threshold=99.9)
    assert response.status_code == 409 and response.json()['detail']['code'] == 'exact_duplicate'
    assert len(list((tmp_path/'data').glob('*.docx'))) == 1


def test_threshold_and_original_not_manual_edits(client):
    did=upload(client, word()).json()['document_id']
    blocks=client.get(f'/api/manage/reports/{did}/contents').json()
    for b in blocks:
        if '设备维修预算' in b['text']:
            client.put('/api/manage/contents/'+b['id'],json={'module':'人力资源','section':'建议','text':'人工维护完全不同的内容'})
    changed = word('新增：建立采购年度审计流程，负责人每月追踪落实情况。')
    blocked=upload(client, changed, title='第二份报告', similarity_threshold=90)
    assert blocked.status_code==409
    detail=blocked.json()['detail']; assert detail['code']=='similar_content'
    assert 90 <= detail['matches'][0]['similarity'] < 99.9
    assert detail['matches'][0]['differences']
    allowed=upload(client, changed, title='第二份报告', similarity_threshold=99.9)
    assert allowed.status_code==201,allowed.text


def test_update_requires_review_and_retains_old_version(client):
    old=upload(client,word()).json()['document_id']
    changed=word('第一版修订：新增后备人才名单。')
    response=upload(client,changed,replace_document_id=old)
    assert response.status_code==409 and response.json()['detail']['code']=='version_review'
    assert len(client.get('/api/status').json()['documents'])==1
    response=upload(client,changed,replace_document_id=old,confirm_update=True)
    assert response.status_code==201,response.text
    new=response.json()['document_id']
    assert [r['id'] for r in client.get('/api/status').json()['documents']]==[new]
    assert client.patch(f'/api/manage/reports/{old}/state',json={'deleted':False}).status_code==409
    third=upload(client,word('第二版修订：调整人才盘点频率。'),replace_document_id=new,confirm_update=True)
    assert third.status_code==201,third.text
    with database.connect() as db:
        assert db.execute('SELECT count(*) n FROM document_versions').fetchone()['n']==2
    exact=upload(client,word(bold=True),title='旧稿再上传')
    assert exact.status_code==409 and exact.json()['detail']['code']=='exact_duplicate'


def test_concurrent_name_collision(client):
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda p:upload(client,p),[word('新增甲事项。'),word('新增乙事项。')]))
    assert sorted(r.status_code for r in results)==[201,409]
    assert len(client.get('/api/status').json()['documents'])==1


@pytest.mark.parametrize('same_bytes', [True, False])
def test_parallel_extraction_rechecks_duplicates_before_publication(client, monkeypatch, same_bytes):
    original=duplicates.extract
    barrier=Barrier(2)
    def overlapping(source):
        result=original(source)
        barrier.wait(timeout=10)
        return result
    monkeypatch.setattr(duplicates,'extract',overlapping)
    first=word()
    second=first if same_bytes else word('新增一次内部复盘。')
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda pair:upload(client,pair[1],title=pair[0]), [('报告甲',first),('报告乙',second)]))
    assert sorted(r.status_code for r in results)==[201,409]
    assert len(client.get('/api/status').json()['documents'])==1
    rejection=next(r.json()['detail']['code'] for r in results if r.status_code==409)
    assert rejection==('exact_duplicate' if same_bytes else 'similar_content')


def test_cached_scores_equal_original_formula_after_corpus_change():
    items=[]
    incoming=duplicates.extract(BytesIO(word('新修订。')))
    for i in range(6):
        profile=duplicates.extract(BytesIO(word(f'企业{i}特有的第{i}号事项。')))
        row={'id':str(i),'title':f'报告{i}','company':'测试企业','deleted':False}
        items.append((row,profile))
        if i>=4:
            parts,weights,unseen=duplicates.weights_for(items)
            expected={row['id']:round(duplicates.similarity(duplicates.shingles(incoming),part,weights,unseen)*100,2)
                      for (row,_),part in zip(items,parts)}
            for _ in range(2):
                actual=duplicates.compare(incoming,items,0,'测试企业')
                assert {r['id']:r['similarity'] for r in actual}==expected


def test_extract_keeps_numeric_differences():
    a=duplicates.extract(BytesIO(word('收益率-1.5%')))
    b=duplicates.extract(BytesIO(word('收益率15%')))
    assert a['text_hash']!=b['text_hash']


def test_table_cell_boundaries_are_not_exact_duplicates():
    profiles=[]
    for values in [('甲乙','丙'),('甲','乙丙')]:
        doc=Document();table=doc.add_table(rows=1,cols=2)
        for cell,value in zip(table.rows[0].cells,values): cell.text=value
        buffer=BytesIO();doc.save(buffer);profiles.append(duplicates.extract(BytesIO(buffer.getvalue())))
    assert profiles[0]['text_hash']==profiles[1]['text_hash']
    assert profiles[0]['tables_hash']!=profiles[1]['tables_hash']
    row={'id':'one','title':'报告','company':'测试企业','deleted':False}
    match=duplicates.compare(profiles[0],[(row,profiles[1])],90,'测试企业')[0]
    assert not match['exact'] and match['tables_changed']


def test_rename_and_invalid_version_cannot_bypass_checks(client):
    old=upload(client,word()).json()['document_id']
    doc=Document();doc.add_paragraph('完全不同的数据集，涉及供应链采购合同核查与交付改善。')
    buffer=BytesIO();doc.save(buffer)
    second=upload(client,buffer.getvalue(),title='供应链报告').json()['document_id']
    result=client.patch(f'/api/manage/reports/{second}/names',json={'company_name':'测试（甲）公司','report_title':'年度报告'})
    assert result.status_code==409
    response=upload(client,word('改动内容'),company='另一家公司',replace_document_id=old,confirm_update=True)
    assert response.status_code==409 and response.json()['detail']['code']=='invalid_version'


def test_template_downweighting():
    common='通用管理制度和政策要求需要贯彻执行。'*100
    items=[({'id':str(i)}, {'paragraphs':[common,f'企业{i}独特的技术设备人员问题'+str(i)*100]}) for i in range(8)]
    parts,weights,unseen=duplicates.weights_for(items)
    assert weights['通用管理制'] < weights['企业0独特']
    assert duplicates.similarity(parts[0],parts[1],weights,unseen)<0.9
