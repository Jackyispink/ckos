"""Regression for saved chapter choices surviving refresh and real DOCX downloads."""
from io import BytesIO

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.industry_research import charts, database, exporter
from app.industry_research.planner import build_plan
from app.industry_research.router import router
from app.industry_research.schemas import ResearchBrief


@pytest.fixture
def chart_report(tmp_path, monkeypatch):
    database.init()
    with database.connect() as db:
        assert db.execute('SELECT current_schema() AS name').fetchone()['name'].startswith('test_ckos_')
    monkeypatch.setattr(charts, 'CHART_DIR', tmp_path / 'charts')
    monkeypatch.setattr(exporter, 'EXPORT_DIR', tmp_path / 'exports')
    brief = ResearchBrief(topic='图表验证')
    pid = database.create_project(brief, build_plan(brief))
    for number in range(1, 11):
        database.update_chapter(pid, number, content=f'第{number}章正文。')
    database.update_chapter(pid, 1, content='2022年市场规模为100亿元。\n2023年市场规模达到150亿元。\n2024年市场规模为210亿元。')
    database.update_chapter(pid, 4, content='2022年销量为200万台。\n2023年销量达到240万台。\n2024年销量为300万台。')
    database.update_chapter(pid, 6, content='平台服务费占收入的30%。')
    with database.connect() as db:
        db.execute("INSERT INTO industry_chapters(project_id,chapter_no,title,questions,queries,content,status) VALUES (%s,0,'摘要','[]','[]','研究摘要','complete')", (pid,))
        db.execute("UPDATE industry_projects SET status='complete' WHERE id=%s", (pid,))
    app = FastAPI(); app.include_router(router)
    with TestClient(app) as client:
        yield client, f'/api/industry/projects/{pid}', pid


def choose(client, base, spec, selected=True, kind=None, title=None):
    return client.put(f'{base}/charts/{spec["id"]}', json={
        'id': spec['id'], 'selected': selected, 'chart_type': kind or spec['chart_types'][0],
        'title': title or spec['title']})


def test_choices_persist_and_only_selected_images_land_in_their_own_chapters(chart_report):
    client, base, pid = chart_report
    candidates = client.get(base + '/charts').json()['charts']
    first = next(s for s in candidates if s['chapter_no'] == 1 and 'line' in s['chart_types'])
    fourth = next(s for s in candidates if s['chapter_no'] == 4 and 'line' in s['chart_types'])
    assert choose(client, base, first, kind='line', title='第一章市场趋势').json()['selected'] is True
    assert choose(client, base, fourth, kind='bar', title='第四章销量对比').json()['selected'] is True
    # Refresh and switch a style: saving one chapter cannot clear the other.
    assert choose(client, base, first, kind='bar', title='第一章市场趋势').status_code == 200
    saved = client.get(base + '/charts').json()
    assert saved['selected_count'] == 2
    response = client.get(base + '/download/docx')
    assert response.status_code == 200, response.text[:300] if response.status_code != 200 else ''
    assert response.headers['x-report-chart-count'] == '2'
    assert response.headers['cache-control'] == 'no-store'
    doc = Document(BytesIO(response.content))
    assert len(doc.inline_shapes) == 2
    chapter = None; locations = []
    for paragraph in doc.paragraphs:
        if paragraph.style.name == 'Heading 1' and paragraph.text.startswith('第'):
            chapter = paragraph.text.split('章')[0]
        for prop in paragraph._p.xpath('.//wp:docPr'):
            locations.append((chapter, prop.get('title')))
    assert locations == [('第1', '第一章市场趋势'), ('第4', '第四章销量对比')]
    # Not drawing is a persisted decision, including all charts disabled.
    assert choose(client, base, first, selected=False).status_code == 200
    assert client.get(base + '/charts').json()['selected_count'] == 1
    second = client.get(base + '/download/docx')
    assert len(Document(BytesIO(second.content)).inline_shapes) == 1
    assert choose(client, base, fourth, selected=False).status_code == 200
    third = client.get(base + '/download/docx')
    assert len(Document(BytesIO(third.content)).inline_shapes) == 0


def test_failed_render_returns_error_instead_of_old_chartless_report(chart_report, monkeypatch):
    client, base, _ = chart_report
    spec = client.get(base + '/charts').json()['charts'][0]
    assert choose(client, base, spec).status_code == 200
    def broken(*args):
        raise RuntimeError('test renderer failure')
    monkeypatch.setattr(charts, 'render', broken)
    response = client.get(base + '/download/docx')
    assert response.status_code == 500
    assert '绘制失败' in response.json()['detail']


def test_changed_data_invalidates_saved_selection_instead_of_silent_omission(chart_report):
    client, base, pid = chart_report
    spec = next(s for s in client.get(base + '/charts').json()['charts'] if s['chapter_no'] == 1)
    assert choose(client, base, spec).status_code == 200
    database.update_chapter(pid, 1, content='本章暂无数据。')
    state = client.get(base + '/charts').json()
    assert state['stale_selected_count'] == 1
    assert client.get(base + '/download/docx').status_code == 409
    assert client.delete(base + '/charts/stale').json()['stale_selected_count'] == 0
    assert client.get(base + '/download/docx').status_code == 200


def test_chart_choice_cannot_target_other_project_or_unsupported_type(chart_report):
    client, base, _ = chart_report
    candidates = client.get(base + '/charts').json()['charts']
    spec = next(s for s in candidates if 'line' in s['chart_types'])
    assert choose(client, base, spec, kind='pie').status_code == 422
    assert client.put(base + '/charts/unknown000', json={'id': 'unknown000', 'title': '伪造图', 'chart_type': 'bar'}).status_code == 422
    assert client.put(base + '/charts', json={'charts': [dict(id=spec['id'], title='重复', chart_type='bar')] * 2}).status_code == 422


def test_single_fact_card_is_exported_without_inventing_a_comparison(chart_report):
    client, base, _ = chart_report
    spec = next(s for s in client.get(base + '/charts').json()['charts'] if s['chapter_no'] == 6)
    assert spec['chart_types'] == ['card']
    assert spec['values'] == [30]
    assert choose(client, base, spec, kind='card').status_code == 200
    response = client.get(base + '/download/docx')
    assert response.status_code == 200
    assert response.headers['x-report-chart-count'] == '1'
    assert len(Document(BytesIO(response.content)).inline_shapes) == 1


@pytest.mark.parametrize('kind', ['column', 'area', 'scatter', 'lollipop'])
def test_new_chart_styles_save_refresh_and_export(chart_report, kind):
    client, base, _ = chart_report
    spec = next(s for s in client.get(base + '/charts').json()['charts'] if s['kind'] == 'time_series')
    assert kind in spec['chart_types']
    assert choose(client, base, spec, kind=kind).status_code == 200
    saved = next(s for s in client.get(base + '/charts').json()['charts'] if s['id'] == spec['id'])
    assert saved['selected_type'] == kind
    response = client.get(base + '/download/docx')
    assert response.status_code == 200
    assert response.headers['x-report-chart-count'] == '1'
    assert len(Document(BytesIO(response.content)).inline_shapes) == 1


def test_extraction_keeps_metrics_separate_and_inherits_short_year_rows():
    project = {'chapters': [{'chapter_no': 4, 'title': '市场', 'content':
        '2024年：规模突破500亿元，增速约34.9%。\n2025年：规模约677.9亿元，增速约34.4%。\n2026年：预计793.3亿元，增速约17.0%。\n2030年：预计达到1505.9亿元。\n2024年制作成本为20亿元。\n2025年制作成本为30亿元。'}]}
    specs = charts.discover(project)
    series = [s for s in specs if s['kind'] == 'time_series']
    assert len(series) == 2
    assert series[0]['values'] == [500, 677.9, 793.3, 1505.9]
    assert series[1]['values'] == [20, 30]
    assert not any('pie' in s['chart_types'] for s in specs)


def test_version_download_and_restore_roundtrip(chart_report):
    client, base, pid = chart_report
    spec = next(s for s in client.get(base + '/charts').json()['charts'] if s['kind'] == 'time_series')
    assert choose(client, base, spec, kind='area').status_code == 200
    original = database.chapter(pid, 1)['content']
    saved = client.post(base + '/versions')
    assert saved.status_code == 201
    vid = saved.json()['id']
    assert len(client.get(base + '/versions').json()) == 1
    database.update_chapter(pid, 1, content='第二份报告的正文。')
    assert choose(client, base, spec, selected=False).status_code == 422  # Data changed.
    download = client.get(base + f'/versions/{vid}/download/docx')
    assert download.status_code == 200
    assert len(Document(BytesIO(download.content)).inline_shapes) == 1
    assert client.post(base + f'/versions/{vid}/restore').status_code == 200
    assert database.chapter(pid, 1)['content'] == original
    assert client.get(base + '/charts').json()['selected_count'] == 1
    rows = client.get(base + '/versions').json()
    assert len(rows) == 2
    other = next(x for x in rows if x['id'] != vid)
    assert client.post(base + f'/versions/{other["id"]}/restore').status_code == 200
    assert database.chapter(pid, 1)['content'] == '第二份报告的正文。'
    assert client.get(base + '/versions/not-found').status_code == 404
